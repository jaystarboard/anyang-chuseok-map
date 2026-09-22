from __future__ import annotations

import json
import os
import re
import sys
from collections import Counter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "data", "facilities.json")
COORDS = os.path.join(ROOT, "tools", "coords.json")


SHEETS = ("동안구", "만안구")
ER_SHEET = "경기도 응급의료기관"
MOON_SHEET = "경기도 달빛어린이병원"
DAYS = [
    {"date": "2026-09-24", "label": "9.24", "dow": "목"},
    {"date": "2026-09-25", "label": "9.25", "dow": "금", "holiday": "추석"},
    {"date": "2026-09-26", "label": "9.26", "dow": "토"},
    {"date": "2026-09-27", "label": "9.27", "dow": "일"},
]
ALWAYS_OPEN = "응급실운영"

CATEGORY = {
    "권역응급의료센터": "응급실", "지역응급의료센터": "응급실",
    "종합병원": "병원", "병원": "병원", "요양병원": "병원",
    "한방병원": "한방", "한의원": "한방",
    "의원": "의원", "치과의원": "치과", "약국": "약국",
}
PHARMACY_KINDS = {"약국"}

ER_TIERS = ("권역센터", "지역센터", "지역기관")


def norm(v) -> str:
    return re.sub(r"\s+", " ", str(v)).strip() if v is not None else ""


def read_sheet(ws, gu: str) -> tuple[list[dict], dict[str, list[int]]]:
    header = next((r for r in range(1, 12)
                   if norm(ws.cell(r, 1).value) == "구분" and norm(ws.cell(r, 2).value) == "명칭"), None)
    if header is None:
        raise SystemExit(f"[{gu}] 헤더 행(구분/명칭)을 찾지 못했습니다. 배포본 양식 변경 확인 필요.")

    rows, subtotals = [], {}
    for r in range(header + 1, ws.max_row + 1):
        kind, name = norm(ws.cell(r, 1).value), norm(ws.cell(r, 2).value)
        if not kind and not name:
            continue
        if kind.endswith("소계") or kind == "소계":
            if name.isdigit():
                subtotals[f"{kind}@{r}"] = [name] + [norm(ws.cell(r, c).value) for c in range(5, 9)]
            continue
        if kind == "구분" or kind.startswith("총 ") or kind.startswith("총계"):
            continue
        if kind not in CATEGORY:
            raise SystemExit(f"[{gu}] r{r} 알 수 없는 구분 {kind!r} — CATEGORY 에 추가하세요.")
        hours = [norm(ws.cell(r, c).value) for c in range(5, 9)]
        if not any(hours):
            raise SystemExit(f"[{gu}] r{r} {name!r} 4일 모두 공란입니다. 양식 확인 필요.")
        rows.append({"gu": gu, "kind": kind, "name": name,
                     "address": norm(ws.cell(r, 3).value), "tel": norm(ws.cell(r, 4).value),
                     "hours": hours})
    return rows, subtotals


def read_xlsx(path: str) -> list[dict]:
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True)
    missing = [s for s in SHEETS if s not in wb.sheetnames]
    if missing:
        raise SystemExit(f"시트 없음: {missing} (있는 시트: {wb.sheetnames})")

    all_rows = []
    for gu in SHEETS:
        rows, subtotals = read_sheet(wb[gu], gu)
        verify(gu, rows, subtotals)
        all_rows += rows
    return all_rows


def verify(gu: str, rows: list[dict], subtotals: dict[str, list[int]]) -> None:
    got_ph = [sum(1 for r in rows if r["kind"] in PHARMACY_KINDS and r["hours"][i]) for i in range(4)]
    got_med = [sum(1 for r in rows if r["kind"] not in PHARMACY_KINDS and r["hours"][i]) for i in range(4)]
    n_ph = sum(1 for r in rows if r["kind"] in PHARMACY_KINDS)
    n_med = len(rows) - n_ph

    want_ph = want_med = None
    for key, vals in subtotals.items():
        nums = [int(v) for v in vals if str(v).isdigit()]
        if len(nums) != 5:
            continue
        total, per_day = nums[0], nums[1:]
        if "약국" in key or (want_ph is None and total == n_ph and per_day == got_ph):
            want_ph = (total, per_day)
        else:
            want_med = (total, per_day) if want_med is None else (want_med[0] + total,
                                                                  [a + b for a, b in zip(want_med[1], per_day)])
    if want_ph is None or want_med is None:
        raise SystemExit(f"[{gu}] 소계 행을 해석하지 못했습니다: {subtotals}")
    if (n_ph, got_ph) != want_ph:
        raise SystemExit(f"[{gu}] 약국 소계 불일치: 추출 {n_ph}{got_ph} vs 시트 {want_ph}")
    if (n_med, got_med) != want_med:
        raise SystemExit(f"[{gu}] 의료기관 소계 불일치: 추출 {n_med}{got_med} vs 시트 {want_med}")
    print(f"  [{gu}] 소계 검증 통과 — 의료기관 {n_med}{got_med}, 약국 {n_ph}{got_ph}")


def merged_value(ws, row: int, col: int):
    v = ws.cell(row, col).value
    if v is not None:
        return v
    for rng in ws.merged_cells.ranges:
        if rng.min_row <= row <= rng.max_row and rng.min_col <= col <= rng.max_col:
            return ws.cell(rng.min_row, rng.min_col).value
    return None


def tel_with_area(tel: str) -> str:
    t = norm(tel).replace(" ", "")
    if not t or t.startswith("0") or re.match(r"^1\d{3}-", t):
        return t
    return f"031-{t}"


def read_er_sheet(wb) -> list[dict]:
    ws = wb[ER_SHEET]
    header = next((r for r in range(1, 6) if norm(ws.cell(r, 1).value) == "시군명"), None)
    if header is None:
        raise SystemExit(f"[{ER_SHEET}] 헤더(시군명)를 찾지 못했습니다.")
    rows = []
    for r in range(header + 1, ws.max_row + 1):
        sigun, name = norm(ws.cell(r, 1).value), norm(ws.cell(r, 2).value)
        tier, addr = norm(ws.cell(r, 3).value), norm(ws.cell(r, 6).value)
        if not name:
            continue
        if tier not in ER_TIERS:
            raise SystemExit(f"[{ER_SHEET}] r{r} 알 수 없는 권역구분명 {tier!r}")
        rows.append({"set": "er", "gu": sigun, "cat": tier, "kind": tier, "name": name,
                     "address": addr, "tel": tel_with_area(ws.cell(r, 4).value),
                     "erTel": tel_with_area(ws.cell(r, 5).value)})
    print(f"  [{ER_SHEET}] {len(rows)}건 — " + ", ".join(
        f"{t} {sum(1 for x in rows if x['cat'] == t)}" for t in ER_TIERS))
    return rows


def read_moon_sheet(wb) -> list[dict]:
    ws = wb[MOON_SHEET]
    rows = []
    for r in range(1, ws.max_row + 1):
        if not isinstance(ws.cell(r, 1).value, (int, float)):
            continue
        name = norm(ws.cell(r, 3).value)
        if not name:
            continue
        weekday = holiday = ""
        for rr in (r, r + 1):
            if rr > ws.max_row or (rr != r and isinstance(ws.cell(rr, 1).value, (int, float))):
                continue
            line = norm(ws.cell(rr, 6).value)
            if not line:
                continue
            head = line.split(")")[0]
            if "평" in head:
                weekday = line
            elif any(k in head for k in ("토", "일", "공")):
                holiday = line
        rows.append({"set": "moon", "gu": norm(merged_value(ws, r, 2)), "cat": "달빛",
                     "kind": "달빛어린이병원", "name": name, "address": "", "tel": "",
                     "partner": norm(ws.cell(r, 4).value),
                     "weekday": weekday, "holiday": holiday})
    no_holiday = [x["name"] for x in rows if not x["holiday"]]
    print(f"  [{MOON_SHEET}] {len(rows)}건" +
          (f" — 휴일시간 미표기 {len(no_holiday)}건: {', '.join(no_holiday)}" if no_holiday else ""))
    return rows


def load_coords() -> dict:
    if not os.path.exists(COORDS):
        raise SystemExit(f"{COORDS} 가 없습니다. tools/geocode.html 을 등록된 도메인에서 실행해 만드세요.")
    raw = json.load(open(COORDS, encoding="utf-8"))
    return {k: v for k, v in raw.items() if not k.startswith("_")}


def coord_key(r: dict) -> str:
    prefix = r.get("set", "anyang")
    return f"{r['gu']}|{r['name']}" if prefix == "anyang" else f"{prefix}|{r['gu']}|{r['name']}"


def attach_coords(rows: list[dict]) -> None:
    coords = load_coords()
    missing = [coord_key(r) for r in rows if coord_key(r) not in coords]
    if missing:
        raise SystemExit("좌표표에 없는 기관:\n  " + "\n  ".join(missing)
                         + "\n-> python tools/make_addresses.py 후 tools/geocode.html 을 실행하세요.")
    for r in rows:
        r["lat"], r["lon"] = coords[coord_key(r)]


def main() -> None:
    if len(sys.argv) < 2:
        raise SystemExit("usage: python tools/build_dataset.py <배포 XLSX 경로>")
    import openpyxl
    wb = openpyxl.load_workbook(sys.argv[1], data_only=True)

    anyang = []
    for gu in SHEETS:
        rows, subtotals = read_sheet(wb[gu], gu)
        verify(gu, rows, subtotals)
        for r in rows:
            r["set"] = "anyang"
        anyang += rows
    er = read_er_sheet(wb)
    moon = read_moon_sheet(wb)

    attach_coords(anyang + er + moon)

    def cat_of(r):
        return "응급실" if ALWAYS_OPEN in r["hours"] else CATEGORY[r["kind"]]

    facilities = []
    for i, r in enumerate(sorted(anyang, key=lambda x: (cat_of(x) != "응급실", x["gu"], cat_of(x), x["name"])), 1):
        facilities.append({"id": i, "gu": r["gu"], "cat": cat_of(r), "kind": r["kind"], "name": r["name"],
                           "addr": r["address"], "tel": r["tel"], "hours": r["hours"],
                           "lat": r["lat"], "lon": r["lon"]})

    er_out = []
    for i, r in enumerate(sorted(er, key=lambda x: (ER_TIERS.index(x["cat"]), x["gu"], x["name"])), 1):
        er_out.append({"id": 1000 + i, "gu": r["gu"], "cat": r["cat"], "kind": r["cat"], "name": r["name"],
                       "addr": r["address"], "tel": r["tel"], "erTel": r["erTel"],
                       "lat": r["lat"], "lon": r["lon"]})

    moon_out = []
    for i, r in enumerate(sorted(moon, key=lambda x: (x["gu"], x["name"])), 1):
        moon_out.append({"id": 2000 + i, "gu": r["gu"], "cat": "달빛", "kind": "달빛어린이병원",
                         "name": r["name"], "addr": "", "tel": "", "partner": r["partner"],
                         "weekday": r["weekday"], "holiday": r["holiday"],
                         "lat": r["lat"], "lon": r["lon"]})

    doc = {
        "meta": {
            "title": "2026 추석 연휴 문여는 병원·약국 (안양시 전체)",
            "source": "안양시보건소 「2026년 추석 연휴 문여는 의료기관 및 약국 현황」 (동안구·만안구)",
            "sourceUrl": "https://www.anyang.go.kr/health/selectBbsNttView.do?key=1369&bbsNo=108&nttNo=458257",
            "sourceAsOf": "2026-09-18",
            "geocoder": "카카오 로컬 주소검색·장소검색 (tools/geocode.html)",
            "contacts": [
                {"label": "응급의료포털 E-Gen", "url": "https://www.e-gen.or.kr"},
                {"label": "보건복지부 콜센터", "tel": "129"},
                {"label": "구급상황관리센터", "tel": "119"},
                {"label": "만안구보건소", "tel": "031-8045-3472"},
                {"label": "동안구보건소", "tel": "031-8045-4472"},
            ],
            "disclaimer": "기관 사정으로 운영시간이 변동될 수 있으니 방문 전 반드시 전화 확인하세요.",
            "days": DAYS,
            "layers": {
                "er": {"label": "경기 응급의료", "note": "24시간 운영. 권역구분명(권역센터·지역센터·지역기관) 기준.",
                       "tiers": list(ER_TIERS)},
                "moon": {"label": "달빛어린이병원",
                         "note": "야간·휴일 소아 진료. 9.24~9.27 은 모두 공휴일이라 (토/일/공) 시간이 적용됩니다."},
            },
        },
        "facilities": facilities,
        "er": er_out,
        "moon": moon_out,
    }
    json.dump(doc, open(OUT, "w", encoding="utf-8"), ensure_ascii=False, indent=1)

    print(f"\n안양시 {len(facilities)} / 경기응급 {len(er_out)} / 달빛 {len(moon_out)} -> {OUT}")
    print("  안양 구:", dict(Counter(f["gu"] for f in facilities)))
    print("  안양 분류:", dict(Counter(f["cat"] for f in facilities)))
    print("  응급 권역구분:", dict(Counter(f["cat"] for f in er_out)))
    for i, d in enumerate(DAYS):
        print(f"  {d['label']}({d['dow']}) 안양 문여는 곳 {sum(1 for f in facilities if f['hours'][i])}곳")


if __name__ == "__main__":
    main()
